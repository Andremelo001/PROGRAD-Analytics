from pathlib import Path

import openpyxl
import pytest


@pytest.fixture
def cpc_xlsx(tmp_path: Path) -> Path:
    """Planilha CPC mínima: header na linha 0, 1 coluna extra, 1 linha inválida."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CPC_2023"
    ws.append(["Ano", "Código da IES", "Nome da IES", "CPC (Faixa)", "Coluna Extra"])
    ws.append([2023, 1, "UF ABC", "4", "-"])
    ws.append([2023, 2, "UF XYZ", "SC", "-"])
    ws.append([2023, "abc", "UF BAD", "3", "-"])  # codigo_ies inválido -> linha dropada
    extra = wb.create_sheet("Atualizações")
    extra.append(["nota de versão"])
    path = tmp_path / "CPC_2023.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def trajetoria_xlsx(tmp_path: Path) -> Path:
    """Planilha Trajetória mínima: 6 linhas de lixo, cabeçalho 6+7, códigos na 8."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "INDICADORES_TRAJETORIA"
    for _ in range(6):
        ws.append([None])
    ws.append(["Código da Instituição", "Ano de Ingresso", "Indicadores de Trajetória"])
    ws.append([None, None, "Taxa de Permanência - TAP"])
    ws.append(["CO_IES", "NU_ANO_INGRESSO", "TAP"])
    ws.append([1, 2019, 98.5])
    ws.append([2, 2020, 91.0])
    path = tmp_path / "trajetoria.xlsx"
    wb.save(path)
    return path
